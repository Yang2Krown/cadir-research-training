#!/usr/bin/env python3
"""Build the report from its Markdown, preserved figures and fixed Git sources."""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.text.paragraph import Paragraph
import fitz

ROOT = Path(__file__).resolve().parents[1]
LM_SHA = 'ab68da7ef154f133ec436275fcf11396f37f884a'
HSC_SOURCE_SHA = '6cc7b3da67c193da11a67c45a75167d6521438b0'
HSC_RESULT_SHA = '091a8a4585cabc59a5504380398892bac49044f4'
REPOSITORY = 'https://github.com/Yang2Krown/cadir-research-training'


def sha(data):
    return hashlib.sha256(data).hexdigest()


def git_bytes(repo, commit, path):
    return subprocess.run(['git', '-C', str(repo), 'show', f'{commit}:{path}'],
                          check=True, capture_output=True).stdout


def collect(repo):
    evidence = ROOT / 'evidence'
    assets = ROOT / 'assets'
    entries = []

    def save(commit, path, target):
        data = git_bytes(repo, commit, path)
        dest = ROOT / target
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(data)
        entries.append(dict(path=target, bytes=len(data), sha256=sha(data),
                            source_commit=commit, source_path=path,
                            source_url=f'{REPOSITORY}/blob/{commit}/{path}'))
        return data

    lm = 'cases/tabula_muris_limb/'
    fig = lm + 'fig3_completion/'
    names = ['Fig3A_association_top5', 'Fig3B_Bcell_principal',
             'Fig3C_endothelial_principal', 'Fig3D_actual_split_merge_history',
             'S1A_direction_angle_heatmap', 'S1B_direction_projection_matrix',
             'S1C_reference_sankey', 'S1D_Bcell_standard', 'S1E_endothelial_standard']
    for name in names:
        save(LM_SHA, fig + f'figures/{name}.png', f'assets/{name}.png')
    for name in ['cluster_annotation_reference_map.tsv', 'all_gene_ranks.tsv',
                 'all_gene_assignments.tsv', 'top5_genes.tsv',
                 'principal_standard_projection_angle_checks.tsv',
                 'Fig3D_history_edges.tsv', 'Fig3D_node_annotations.tsv']:
        save(LM_SHA, fig + 'tables/' + name, 'evidence/LM/' + name)
    for name in ['metrics.json', 'selection.json', 'ranking_provenance.json',
                 'stage_times_seconds.json', 'required_version_check.tsv']:
        save(LM_SHA, fig + 'evidence/real-run/' + name, 'evidence/LM/' + name)
    for name in ['run_config.json', 'real_run.json']:
        save(LM_SHA, fig + 'evidence/' + name, 'evidence/LM/' + name)
    hsc = 'cases/nestorowa_hsc/fair-experiment/'
    result = hsc + 'results/6cc7b3d-20261027-attempt-01/'
    for name in ['04-common-cohort.png', '01-all-seed-ARI-NMI.png',
                 '02-cadir-seed2358-reference-share.png',
                 '02-pca_snn_louvain-seed2358-reference-share.png']:
        save(HSC_RESULT_SHA, result + 'presentation/' + name, 'assets/HSC_' + name)
    for name in ['all-seed-presented-values.tsv', 'cohort-summary.tsv']:
        save(HSC_RESULT_SHA, result + 'presentation/' + name, 'evidence/HSC/' + name)
    save(HSC_RESULT_SHA, result + 'formal/cadir-seed-2358-attempt-01/sessionInfo.txt',
         'evidence/HSC/sessionInfo.txt')
    source_specs = []
    for name in ['pipeline.R', 'run_case.py', 'rerun.sh', 'run-r.sh', 'verify_results.R']:
        source_specs.append((LM_SHA, lm + 'scripts/' + name, 'evidence/source/LM/' + name))
    for name in ['render_fig3.R', 'repair_s1_legend.R']:
        source_specs.append((LM_SHA, fig + 'scripts/' + name, 'evidence/source/LM/' + name))
    for name in ['prepare_common.R', 'formal_run.R', 'run_experiment.py', 'rank_diagnostics.R']:
        source_specs.append((HSC_SOURCE_SHA, hsc + 'scripts/' + name, 'evidence/source/HSC/' + name))
    for spec in source_specs:
        save(*spec)
    save(HSC_SOURCE_SHA, hsc + 'design/scientific-config.json', 'evidence/HSC/scientific-config.json')
    with (evidence / '代码段索引.tsv').open('w') as f:
        writer = csv.writer(f, delimiter='\t', lineterminator='\n')
        writer.writerow(['module', 'fixed_commit', 'source_path', 'start_line', 'excerpt_path'])
        for commit, source, target in source_specs:
            data = (ROOT / target).read_text()
            patterns = ['audit_object <-', 'export_results <-', 'modelGeneVar',
                        'do.call(APL::cacomp', 'CAdir::dirclust_splitmerge',
                        'CAdir::rank_genes', 'aricode::clustComp', 'CAdir::annotate_biclusters',
                        'scater::runPCA', 'scran::buildSNNGraph', 'igraph::cluster_louvain',
                        'scuttle::logNormCounts', 'def main', 'tempfile.mkdtemp']
            for pattern in patterns:
                for n, line in enumerate(data.splitlines(), 1):
                    if pattern in line:
                        writer.writerow([pattern, commit, source, n, target])
    manifest = dict(schema_version=1, report='实验报告.md',
                    management_base='8d86e93104bae3f01c7b15bc0b07b811b62884f6',
                    collection='byte-exact Git object reads; no scientific rerun', files=entries)
    (evidence / '材料来源索引.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')


def make_reference(path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)
    section.header_distance, section.footer_distance = Cm(1.2), Cm(1.2)
    for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact', 'Caption']:
        if name not in doc.styles:
            from docx.enum.style import WD_STYLE_TYPE
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(10.5 if name != 'Caption' else 9)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2
    for name, size in [('Title', 22), ('Subtitle', 12), ('Heading 1', 14),
                       ('Heading 2', 10.5), ('Heading 3', 10.5)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        style.paragraph_format.keep_with_next = True
    for name in ['Source Code', 'Verbatim Char']:
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = 'Liberation Mono'
        style.font.size = Pt(8)
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('CAdir单细胞联合聚类复现与新数据对照实验').font.size = Pt(9)
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('— ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    p._p.append(field)
    p.add_run(' —')
    doc.save(path)


def insert_before(anchor, owner, text='', style=None):
    el = OxmlElement('w:p')
    anchor.addprevious(el)
    p = Paragraph(el, owner)
    if style:
        p.style = style
    p.add_run(text)
    return p


def build():
    reference = ROOT / 'scripts/reference.docx'
    make_reference(reference)
    subprocess.run(['pandoc', '实验报告.md', '--from=markdown+tex_math_dollars+implicit_figures',
                    '--standalone', '--toc', '--toc-depth=2',
                    '--reference-doc=' + str(reference), '--resource-path=.',
                    '--output=实验报告.docx'], cwd=ROOT, check=True)
    path = ROOT / '实验报告.docx'
    doc = Document(path)
    headings = [(p.text, p.style.name) for p in doc.paragraphs
                if p.style.name in ['Heading 1', 'Heading 2']]
    sdt = next(el for el in doc.element.body if el.tag == qn('w:sdt'))
    toc_title = insert_before(sdt, doc._body, '目录', 'Heading 1')
    toc_title.paragraph_format.page_break_before = True
    toc_entries = []
    for text, style in headings:
        p = insert_before(sdt, doc._body, text + '\t000', 'Normal')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.left_indent = Cm(.5 if style == 'Heading 2' else 0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.8),
                                                WD_TAB_ALIGNMENT.RIGHT,
                                                WD_TAB_LEADER.DOTS)
        for r in p.runs:
            r.font.size = Pt(9)
        toc_entries.append((p, text))
    sdt.getparent().remove(sdt)
    for p in doc.paragraphs:
        if p.text == '摘要':
            p.paragraph_format.page_break_before = True
            break
    # Keep table rows intact and make the first row repeat on page breaks.
    for table in doc.tables:
        table.autofit = False
        ncols = len(table.columns)
        for col in table.columns:
            col.width = Cm(16 / ncols)
        for i, row in enumerate(table.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if i == 0:
                pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        run.font.size = Pt(8.5)
    for p in doc.paragraphs:
        if p.style.name == 'Source Code':
            p.paragraph_format.keep_together = False
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.size = Pt(8)
        if p.style.name in ['Caption', 'Image Caption']:
            p.paragraph_format.keep_together = True
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = 'CAdir单细胞联合聚类的复现与新数据对照实验'
    doc.core_properties.subject = '科研基础训练实验报告'
    doc.core_properties.author = ''
    doc.core_properties.comments = ''
    doc.save(path)
    soffice = shutil.which('soffice')
    def export_pdf():
        proc = subprocess.run([soffice, '-env:UserInstallation=file:///tmp/cadir-report-lo-profile',
                              '--headless', '--convert-to', 'pdf', '--outdir', str(ROOT), str(path)],
                             check=True, timeout=120, capture_output=True)
        if not (ROOT / '实验报告.pdf').exists():
            raise RuntimeError(proc.stdout.decode(errors='replace') + proc.stderr.decode(errors='replace'))
    export_pdf()
    pdf = fitz.open(ROOT / '实验报告.pdf')
    norm = lambda s: ''.join(s.split())
    page_map = {norm(title): page for _, title, page in pdf.get_toc()}
    pdf.close()
    for p, title in toc_entries:
        number = page_map[norm(title)]
        p.runs[0].text = title + '\t' + str(number)
    doc.save(path)
    export_pdf()
    pdf = fitz.open(ROOT / '实验报告.pdf')
    final_map = {norm(title): page for _, title, page in pdf.get_toc()}
    assert all(page_map[norm(t)] == final_map[norm(t)] for _, t in toc_entries)
    (ROOT / 'evidence/目录页码.json').write_text(json.dumps(
        {t: final_map[norm(t)] for _, t in toc_entries}, ensure_ascii=False, indent=2) + '\n')
    pdf.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source-repo', type=Path)
    parser.add_argument('--collect-only', action='store_true')
    args = parser.parse_args()
    if args.source_repo:
        collect(args.source_repo)
    if not args.collect_only:
        build()
    print(json.dumps({'markdown_bytes': (ROOT / '实验报告.md').stat().st_size,
                      'figures': len(list((ROOT / 'assets').glob('*.png')))}, ensure_ascii=False))


if __name__ == '__main__':
    main()
